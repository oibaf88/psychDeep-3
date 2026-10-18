"""Private clinical-document upload and whole-person memory controls."""
from __future__ import annotations

import hashlib
import io
import re
import uuid
import zipfile
from datetime import datetime, timezone
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import PatientProfessionalAssignment, User
from app.models_vnext import PatientClinicalDocument
from app.security import require_professional
from app.services import audit
from app.services.consent import CLINICAL_DOCUMENT_ANALYSIS, LINGUISTIC_ANALYSIS, require_granted
from app.services.memory_pipeline import enqueue_memory_job

router = APIRouter(prefix="/api/v1/professional/patients", tags=["clinical-documents"])

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_EXTRACTED_CHARS = 250_000
MAX_PDF_PAGES = 250
MAX_DOCX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_DOCX_ENTRIES = 512
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain; charset=utf-8",
}


class DocumentMetadata(BaseModel):
    id: uuid.UUID
    original_filename: str
    media_type: str
    file_size_bytes: int
    text_truncated: bool
    extraction_status: str
    analysis_status: str
    error_category: str | None
    patient_model_version_id: uuid.UUID | None
    uploaded_at: datetime
    analyzed_at: datetime | None


class MemoryRebuildAccepted(BaseModel):
    status: str
    message: str


def _require_clinical_read(db: Session, professional: User, patient_id: uuid.UUID) -> None:
    patient_exists = db.query(User.id).filter(User.id == patient_id, User.role == "patient").first()
    if patient_exists is None:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")
    if professional.role == "admin_clinical":
        raise HTTPException(status_code=403, detail="admin_clinical no tiene visibilidad clínica de documentos")
    if professional.role == "therapist":
        assigned = (
            db.query(PatientProfessionalAssignment.id)
            .filter(
                PatientProfessionalAssignment.patient_id == patient_id,
                PatientProfessionalAssignment.professional_id == professional.id,
                PatientProfessionalAssignment.status.in_(("active", "paused")),
            )
            .first()
        )
        if assigned is None:
            raise HTTPException(status_code=403, detail="No tienes asignación activa/pausada con este paciente")
        return
    if professional.role != "supervisor":
        raise HTTPException(status_code=403, detail="No tienes acceso a documentos clínicos")


def _require_active_therapist(db: Session, professional: User, patient_id: uuid.UUID) -> None:
    if professional.role != "therapist":
        raise HTTPException(status_code=403, detail="Solo el terapeuta asignado puede incorporar documentos")
    patient_exists = db.query(User.id).filter(User.id == patient_id, User.role == "patient").first()
    if patient_exists is None:
        raise HTTPException(status_code=404, detail="Paciente no encontrado")
    assigned = (
        db.query(PatientProfessionalAssignment.id)
        .filter(
            PatientProfessionalAssignment.patient_id == patient_id,
            PatientProfessionalAssignment.professional_id == professional.id,
            PatientProfessionalAssignment.status == "active",
        )
        .first()
    )
    if assigned is None:
        raise HTTPException(status_code=403, detail="Se requiere asignación activa para incorporar documentos")


def _safe_filename(value: str | None) -> str:
    name = (value or "informe").replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return (name or "informe")[:240]


def _extract_text(extension: str, content: bytes) -> str:
    if extension == ".txt":
        extracted = content.decode("utf-8-sig", errors="replace")
    elif extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("invalid_pdf")
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ValueError("encrypted_pdf")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("pdf_page_limit")
        pages: list[str] = []
        total_chars = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)
            total_chars += len(page_text)
            if total_chars >= MAX_EXTRACTED_CHARS:
                break
        extracted = "\n\n".join(pages)
    elif extension == ".docx":
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise ValueError("invalid_docx")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_ENTRIES or sum(member.file_size for member in members) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("docx_expansion_limit")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ValueError("invalid_docx")
        from docx import Document

        doc = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]
        for table in doc.tables:
            for row in table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        extracted = "\n".join(parts)
    else:
        raise ValueError("unsupported_document_type")

    extracted = extracted.replace("\r\n", "\n").replace("\r", "\n")
    extracted = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", extracted).strip()
    return extracted


def _metadata_row(row) -> dict:
    return {
        "id": row.id,
        "original_filename": row.original_filename,
        "media_type": row.media_type,
        "file_size_bytes": row.file_size_bytes,
        "text_truncated": row.text_truncated,
        "extraction_status": row.extraction_status,
        "analysis_status": row.analysis_status,
        "error_category": row.error_category,
        "patient_model_version_id": row.patient_model_version_id,
        "uploaded_at": row.uploaded_at,
        "analyzed_at": row.analyzed_at,
    }


@router.get("/{patient_id}/documents", response_model=list[DocumentMetadata])
def list_patient_documents(
    patient_id: uuid.UUID,
    db: Session = Depends(get_db),
    professional: User = Depends(require_professional),
):
    _require_clinical_read(db, professional, patient_id)
    rows = (
        db.query(
            PatientClinicalDocument.id,
            PatientClinicalDocument.original_filename,
            PatientClinicalDocument.media_type,
            PatientClinicalDocument.file_size_bytes,
            PatientClinicalDocument.text_truncated,
            PatientClinicalDocument.extraction_status,
            PatientClinicalDocument.analysis_status,
            PatientClinicalDocument.error_category,
            PatientClinicalDocument.patient_model_version_id,
            PatientClinicalDocument.uploaded_at,
            PatientClinicalDocument.analyzed_at,
        )
        .filter(PatientClinicalDocument.user_id == patient_id)
        .order_by(PatientClinicalDocument.uploaded_at.desc())
        .limit(100)
        .all()
    )
    audit.log(
        db,
        actor_id=professional.id,
        actor_role=professional.role,
        action="patient_clinical_documents_viewed",
        entity_type="patient_clinical_documents",
        entity_id=patient_id,
        extra={"patient_id": str(patient_id), "row_count": len(rows)},
    )
    return [_metadata_row(row) for row in rows]


@router.post("/{patient_id}/documents", response_model=DocumentMetadata, status_code=201)
async def upload_patient_document(
    patient_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    professional: User = Depends(require_professional),
):
    _require_active_therapist(db, professional, patient_id)
    require_granted(db, patient_id, LINGUISTIC_ANALYSIS)
    require_granted(db, patient_id, CLINICAL_DOCUMENT_ANALYSIS)

    filename = _safe_filename(file.filename)
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    media_type = MEDIA_TYPES.get(extension)
    if not media_type:
        raise HTTPException(status_code=415, detail="Formato no admitido. Sube un PDF, DOCX o TXT.")

    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if not content:
        raise HTTPException(status_code=422, detail="El archivo está vacío.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="El tamaño máximo por documento es 8 MiB.")

    try:
        extracted = _extract_text(extension, content)
    except Exception as exc:  # noqa: BLE001
        if type(exc).__name__ in {"PdfReadError", "BadZipFile", "PackageNotFoundError"} or isinstance(exc, ValueError):
            raise HTTPException(status_code=422, detail="No se pudo leer el documento. Comprueba el archivo y vuelve a intentarlo.") from None
        raise HTTPException(status_code=422, detail="No se pudo extraer texto del documento.") from None

    truncated = len(extracted) > MAX_EXTRACTED_CHARS
    extracted = extracted[:MAX_EXTRACTED_CHARS]
    extraction_failed = not extracted.strip()
    digest = hashlib.sha256(content).hexdigest()
    doc = PatientClinicalDocument(
        id=uuid.uuid4(),
        user_id=patient_id,
        uploaded_by=professional.id,
        original_filename=filename,
        media_type=media_type,
        file_size_bytes=len(content),
        file_sha256=digest,
        file_content=content,
        extracted_text=extracted,
        text_truncated=truncated,
        extraction_status="failed" if extraction_failed else "text_extracted",
        analysis_status="failed" if extraction_failed else "queued",
        error_category="no_extractable_text" if extraction_failed else None,
    )
    document_id = doc.id
    db.add(doc)
    try:
        if not extraction_failed:
            # Recheck at the write boundary. A revoked grant leaves the file
            # with the uploader and prevents creating a processing job.
            require_granted(db, patient_id, LINGUISTIC_ANALYSIS)
            require_granted(db, patient_id, CLINICAL_DOCUMENT_ANALYSIS)
            enqueue_memory_job(
                db,
                user_id=patient_id,
                job_type="longitudinal_analysis",
                source_table="patient_clinical_documents",
                source_id=doc.id,
                idempotency_key=f"clinical_document:{doc.id}",
            )
        audit.log(
            db,
            actor_id=professional.id,
            actor_role=professional.role,
            action="patient_clinical_document_uploaded",
            entity_type="patient_clinical_document",
            entity_id=doc.id,
            extra={
                "patient_id": str(patient_id),
                "file_size_bytes": len(content),
                "media_type": media_type,
                "analysis_queued": not extraction_failed,
            },
        )
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Ese documento ya se había subido para este paciente.") from None
    except Exception:
        db.rollback()
        raise
    saved = (
        db.query(
            PatientClinicalDocument.id,
            PatientClinicalDocument.original_filename,
            PatientClinicalDocument.media_type,
            PatientClinicalDocument.file_size_bytes,
            PatientClinicalDocument.text_truncated,
            PatientClinicalDocument.extraction_status,
            PatientClinicalDocument.analysis_status,
            PatientClinicalDocument.error_category,
            PatientClinicalDocument.patient_model_version_id,
            PatientClinicalDocument.uploaded_at,
            PatientClinicalDocument.analyzed_at,
        )
        .filter(PatientClinicalDocument.id == document_id)
        .one()
    )
    return _metadata_row(saved)


@router.get("/{patient_id}/documents/{document_id}/download")
def download_patient_document(
    patient_id: uuid.UUID,
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    professional: User = Depends(require_professional),
):
    _require_clinical_read(db, professional, patient_id)
    row = (
        db.query(
            PatientClinicalDocument.original_filename,
            PatientClinicalDocument.media_type,
            PatientClinicalDocument.file_content,
        )
        .filter(
            PatientClinicalDocument.id == document_id,
            PatientClinicalDocument.user_id == patient_id,
        )
        .first()
    )
    if row is None:
        raise HTTPException(status_code=404, detail="No se encontró el documento.")
    audit.log(
        db,
        actor_id=professional.id,
        actor_role=professional.role,
        action="patient_clinical_document_downloaded",
        entity_type="patient_clinical_document",
        entity_id=document_id,
        extra={"patient_id": str(patient_id)},
    )
    encoded_name = quote(row.original_filename, safe="")
    return StreamingResponse(
        io.BytesIO(row.file_content),
        media_type=row.media_type,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}",
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "private, no-store",
        },
    )


@router.post("/{patient_id}/memory/rebuild", response_model=MemoryRebuildAccepted, status_code=202)
def request_patient_model_rebuild(
    patient_id: uuid.UUID,
    db: Session = Depends(get_db),
    professional: User = Depends(require_professional),
):
    _require_active_therapist(db, professional, patient_id)
    require_granted(db, patient_id, LINGUISTIC_ANALYSIS)
    active = db.execute(
        text("""
            select status
              from psychdeep_v12.memory_jobs
             where user_id = :user_id
               and job_type = 'longitudinal_analysis'
               and status in ('pending', 'retry', 'processing')
             order by created_at desc
             limit 1
        """),
        {"user_id": str(patient_id)},
    ).mappings().first()
    if active is not None:
        return MemoryRebuildAccepted(status=str(active["status"]), message="El análisis longitudinal ya está en curso.")

    minute_key = datetime.now(timezone.utc).strftime("%Y%m%d%H%M")
    enqueue_memory_job(
        db,
        user_id=patient_id,
        job_type="longitudinal_analysis",
        source_table="patient_model",
        source_id=patient_id,
        idempotency_key=f"professional_rebuild:{minute_key}",
    )
    audit.log(
        db,
        actor_id=professional.id,
        actor_role=professional.role,
        action="patient_model_rebuild_queued",
        entity_type="patient_model",
        entity_id=patient_id,
        extra={"patient_id": str(patient_id)},
    )
    return MemoryRebuildAccepted(status="queued", message="El análisis se ha añadido a la cola de memoria.")
